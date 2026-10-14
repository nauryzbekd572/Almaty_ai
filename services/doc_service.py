import os
from docx import Document
from docx.shared import Pt, Inches

class DocGenerator:
    @staticmethod
    def generate_complaint_docx(user_name: str, issue_type: str, text: str, ai_law_ref: str, filename: str) -> str:
        """
        Генерирует юридически оформленное заявление в формате .docx
        """
        doc = Document()
        
        # Настройка полей (профессиональный вид)
        for section in doc.sections:
            section.top_margin = Inches(0.8)
            section.bottom_margin = Inches(0.8)
            section.left_margin = Inches(1.0)
            section.right_margin = Inches(0.6)

        # Стилизация текста под Times New Roman (стандарт РК, см. скриншот 1) [5q32js]
        style = doc.styles['Normal']
        font = style.font
        font.name = 'Times New Roman'
        font.size = Pt(12)

        # "Шапка" заявления в правом верхнем углу
        p_header = doc.add_paragraph()
        p_header.alignment = 2  # Правый край
        p_header.add_run("В уполномоченные органы Республики Казахстан\n").bold = True
        p_header.add_run("В рамках общественной инициативы «Адал қоғам»\n")
        p_header.add_run(f"От гражданина(ки): {user_name}\n")
        
        doc.add_paragraph("\n\n")
        
        # Заголовок
        title = doc.add_paragraph()
        title.alignment = 1  # Центр
        title.add_run(f"ЗАЯВЛЕНИЕ\n(По направлению: {issue_type})").bold = True
        
        # Суть обращения
        doc.add_heading("1. Описание инцидента / факта правонарушения:", level=2)
        doc.add_paragraph(text)
        
        # Правовое обоснование от ИИ
        doc.add_heading("2. Правовое обоснование (Конституция 2026 года):", level=2)
        doc.add_paragraph(ai_law_ref)
        
        # Подвал
        doc.add_paragraph("\n\n")
        footer = doc.add_paragraph()
        footer.add_run("Дата подачи: _________________             Подпись: _________________")
        
        # Сохраняем
        doc.save(filename)
        return filename
